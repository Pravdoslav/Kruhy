"""Import the Kruhy Word manuscript into editable Quarto chapters.

The Word file remains the archival source.  This importer transfers the
structure that matters on the web: headings, ordered and nested lists,
hyperlinks, cross-references and tables.  Run it only when intentionally
refreshing the generated chapter text, as it overwrites the chapter files.
"""

from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source" / "Kruhy-1.0d.docx"
BOOK = ROOT / "book"
ONE_SENTENCE_PARTIAL = BOOK / "_myslienka-jednou-vetou.qmd"

PUBLIC_TEXT_REPLACEMENTS = {
    "Ing. Marián Pravdoslav Orávik": "Pravdoslav",
}

CHAPTERS = {
    "Úvodom?": "02-uvodom.qmd",
    "Základy?": "03-zaklady.qmd",
    "Aké sú princípy Kruhov?": "04-principy-kruhov.qmd",
    "Aké sú ďalšie odporúčania pre členov Kruhov?": "05-dalsie-odporucania.qmd",
    "Ako prepojiť, rozšíriť a uplatniť kruhy?": "06-prepojit-rozsirit-uplatnit.qmd",
    "Prečo Celoslovenské Kruhy?": "07-celoslovenske-kruhy.qmd",
    "A čo ďalej? Je toto všetko? Nedá sa štát organizovať nejako lepšie?": "08-a-co-dalej.qmd",
    "Dodatky a diskusia?": "09-dodatky-a-diskusia.qmd",
}


def body_items(doc: Document) -> Iterator[Paragraph | Table]:
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def plain_text(paragraph: Paragraph) -> str:
    text = paragraph.text.replace("\u00a0", " ").strip()
    for original, replacement in PUBLIC_TEXT_REPLACEMENTS.items():
        text = text.replace(original, replacement)
    return text


def paragraph_is_italic(paragraph: Paragraph) -> bool:
    """Return True when every visible Word run is explicitly italic."""
    visible_runs = []
    for run in paragraph._p.xpath(".//w:r"):
        run_text = "".join(node.text or "" for node in run.findall(".//" + qn("w:t")))
        if run_text.strip():
            visible_runs.append(run)
    if not visible_runs:
        return False
    for run in visible_runs:
        properties = run.find(qn("w:rPr"))
        italic = properties.find(qn("w:i")) if properties is not None else None
        if italic is None or italic.get(qn("w:val"), "1") in {"0", "false", "off"}:
            return False
    return True


def italic_markdown(text: str) -> str:
    """Wrap every visible line so Word line breaks do not break Markdown emphasis."""
    return "\n".join(f"*{line}*" if line else "" for line in text.split("\n"))


def paragraph_is_indented_note(paragraph: Paragraph) -> bool:
    """Match italic normal paragraphs indented about 1.25 cm in Word."""
    indent = paragraph.paragraph_format.left_indent
    return (
        paragraph.style.name == "Normal"
        and paragraph_is_italic(paragraph)
        and indent is not None
        and indent.cm >= 1.0
    )


def extract_one_sentence_partial(output: dict[str, list[str]]) -> None:
    """Keep the one-sentence summary reusable on the home page and in chapter 3."""
    chapter = output["03-zaklady.qmd"]
    heading = "## Myšlienka Kruhov 1 vetou?"
    start = chapter.index(heading)
    end = next(
        (index for index in range(start + 1, len(chapter)) if chapter[index].startswith("## ")),
        len(chapter),
    )
    body = chapter[start + 2 : end]
    while body and not body[-1]:
        body.pop()
    ONE_SENTENCE_PARTIAL.write_text("\n".join(body).rstrip() + "\n", encoding="utf-8")
    chapter[start + 2 : end] = [
        "{{< include book/_myslienka-jednou-vetou.qmd >}}",
        "",
    ]


def bookmark_slug(name: str) -> str:
    text = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    anchor_id = f"word-{text or 'bookmark'}"
    aliases = {
        "word-vysvetlenie-preco-su": "word-preco-su-potrebni",
        "word-aky-je-optimalny": "word-optimalny-pocet-clenov",
        "word-aky-je-rozumny": "word-optimalny-pocet-clenov",
    }
    return aliases.get(anchor_id, anchor_id)


def list_info(paragraph: Paragraph, formats: dict[tuple[int, int], str]) -> tuple[int, bool] | None:
    ppr = paragraph._p.pPr
    if ppr is None or ppr.numPr is None:
        return None
    num_id = int(ppr.numPr.numId.val)
    level = int(ppr.numPr.ilvl.val) if ppr.numPr.ilvl is not None else 0
    number_format = formats.get((num_id, level), "bullet")
    return level, number_format != "bullet"


def numbering_formats(doc: Document) -> dict[tuple[int, int], str]:
    """Return the Word number format for each (numId, nesting level)."""
    root = doc.part.numbering_part.element
    abstract_by_num: dict[int, int] = {}
    for num in root.findall(qn("w:num")):
        abstract = num.find(qn("w:abstractNumId"))
        if abstract is not None:
            abstract_by_num[int(num.get(qn("w:numId")))] = int(abstract.get(qn("w:val")))

    formats: dict[tuple[int, int], str] = {}
    for abstract in root.findall(qn("w:abstractNum")):
        abstract_id = int(abstract.get(qn("w:abstractNumId")))
        num_ids = [num_id for num_id, value in abstract_by_num.items() if value == abstract_id]
        for level in abstract.findall(qn("w:lvl")):
            level_number = int(level.get(qn("w:ilvl")))
            fmt = level.find(qn("w:numFmt"))
            value = fmt.get(qn("w:val")) if fmt is not None else "bullet"
            for num_id in num_ids:
                formats[(num_id, level_number)] = value
    return formats


def hyperlinks(paragraph: Paragraph, target_map: dict[str, tuple[str, str]]) -> str:
    """Add Markdown links without touching ordinary Word character styling."""
    text = plain_text(paragraph)
    if not text:
        return text
    fragments: list[tuple[str, str | None]] = []
    for link in paragraph._p.findall(qn("w:hyperlink")):
        label = "".join(node.text or "" for node in link.findall(".//" + qn("w:t"))).replace("\u00a0", " ")
        relation_id = link.get(qn("r:id"))
        anchor = link.get(qn("w:anchor"))
        target: str | None = None
        if relation_id and relation_id in paragraph.part.rels:
            target = paragraph.part.rels[relation_id].target_ref
        elif anchor and anchor in target_map:
            _, target_id = target_map[anchor]
            target = f"#{target_id}"
        if label and target:
            fragments.append((label, target))

    if not fragments:
        return text
    rendered: list[str] = []
    cursor = 0
    for label, target in fragments:
        position = text.find(label, cursor)
        if position < 0:
            continue
        rendered.append(text[cursor:position])
        rendered.append(f"[{label}]({target})")
        cursor = position + len(label)
    rendered.append(text[cursor:])
    return "".join(rendered)


def render_table(table: Table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.replace("\u00a0", " ").split()).replace("|", "\\|") for cell in row.cells]
        rows.append(cells)
    if not rows:
        return []
    width = len(rows[0])
    output = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    output.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return output


def bookmark_lines(paragraph: Paragraph, target_map: dict[str, tuple[str, str]]) -> list[str]:
    """Emit stable HTML anchors only for bookmarks that are linked elsewhere."""
    ids = []
    for bookmark in paragraph._p.findall(".//" + qn("w:bookmarkStart")):
        name = bookmark.get(qn("w:name"))
        if name in target_map:
            ids.append(target_map[name][1])
    return [f'[]{{#{anchor_id}}}' for anchor_id in dict.fromkeys(ids)]


def main() -> None:
    doc = Document(SOURCE)
    formats = numbering_formats(doc)
    target_map: dict[str, tuple[str, str]] = {}
    current = ""
    referenced_anchors = {
        link.get(qn("w:anchor"))
        for paragraph in doc.paragraphs
        for link in paragraph._p.findall(qn("w:hyperlink"))
        if link.get(qn("w:anchor"))
    }
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1" and plain_text(paragraph) in CHAPTERS:
            current = CHAPTERS[plain_text(paragraph)]
        for bookmark in paragraph._p.findall(".//" + qn("w:bookmarkStart")):
            name = bookmark.get(qn("w:name"))
            if name in referenced_anchors:
                target_map[name] = (current, bookmark_slug(name))

    output: dict[str, list[str]] = defaultdict(list)
    current = None
    source_paragraphs = imported_paragraphs = 0
    drawing_number = 0
    open_list = False
    for item in body_items(doc):
        if isinstance(item, Table):
            if current is None:
                raise ValueError("Table found before the first chapter heading")
            if open_list:
                output[current].append("")
                open_list = False
            output[current].extend(render_table(item) + [""])
            continue

        has_drawing = bool(item._p.xpath(".//w:drawing") or item._p.xpath(".//w:pict"))
        text = hyperlinks(item, target_map)
        if text and paragraph_is_italic(item):
            text = italic_markdown(text)
        marks = bookmark_lines(item, target_map)
        if not text and not has_drawing:
            continue
        if text:
            source_paragraphs += 1
        style = item.style.name
        if style == "Title":
            continue
        if style == "Heading 1":
            if open_list:
                output[current].append("")
                open_list = False
            try:
                current = CHAPTERS[plain_text(item)]
            except KeyError as error:
                raise ValueError(f"Unexpected top-level heading: {plain_text(item)}") from error
            output[current].extend(marks + ([""] if marks else []) + [f"# {text}", ""])
        elif current is not None and style.startswith("Heading "):
            if open_list:
                output[current].append("")
                open_list = False
            level = int(style.removeprefix("Heading "))
            output[current].extend(marks + ([""] if marks else []) + [f"{'#' * level} {text}", ""])
            imported_paragraphs += 1
        elif current is not None and has_drawing:
            if open_list:
                output[current].append("")
                open_list = False
            drawing_number += 1
            output[current].extend(marks + [
                f"![Schéma Kruhov – obrázok {drawing_number}](images/inkscape/obr-{drawing_number:02d}.svg)",
                "",
            ])
        elif current is not None:
            info = list_info(item, formats)
            if info:
                level, ordered = info
                marker = "1. " if ordered else "- "
                if marks:
                    if open_list:
                        output[current].append("")
                    output[current].extend(marks)
                output[current].append("    " * level + marker + text)
                open_list = True
            else:
                if open_list:
                    output[current].append("")
                    open_list = False
                if paragraph_is_indented_note(item):
                    output[current].extend(marks + ["::: {.word-note}", text, ":::", ""])
                else:
                    output[current].extend(marks + [text, ""])
            imported_paragraphs += 1

    missing = set(CHAPTERS.values()) - set(output)
    if missing:
        raise ValueError(f"Missing chapters: {', '.join(sorted(missing))}")
    extract_one_sentence_partial(output)
    for filename, content in output.items():
        (BOOK / filename).write_text("\n".join(content).rstrip() + "\n", encoding="utf-8")
    print(f"Imported {imported_paragraphs} of {source_paragraphs} non-empty source paragraphs into {len(output)} chapters.")


if __name__ == "__main__":
    main()
